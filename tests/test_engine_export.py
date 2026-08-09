from itertools import product
from pathlib import Path
import unittest

from openpyxl import load_workbook

from qasas.engine import analyse, bounded_levenshtein
from qasas.export import export_result_xlsx
from qasas.models import DatabaseData, DatabaseEntry, SampleClone, SampleData
from qasas.modes import ALGORITHM_VERSION, MatchingMode
from tests.support import test_path


class EngineExportTests(unittest.TestCase):
    def _sample(self) -> SampleData:
        clones = (
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABCDE", 10, 50.0),
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABCXE", 5, 25.0),
            SampleClone(("IGHV1-2",), ("IGHJ4",), "ABXYE", 3, 15.0),
            SampleClone(("IGHV3-23",), ("IGHJ6",), "ABCDE", 2, 10.0),
        )
        return SampleData(Path("sample.csv"), "CPM", "S1", clones, 20, 4, 4, 0)

    def _database(self) -> DatabaseData:
        entries = (
            DatabaseEntry("IGHV1-2", "IGHJ4", "ABCDE", {"Name": ("KnownAb",)}),
            DatabaseEntry("IGHV3-23", "IGHJ4", "ABCDE", {"Name": ("WrongJ",)}),
        )
        return DatabaseData(
            Path("db.csv"), entries, 2, 2, 0, ("Name",), "Heavy V Gene", "Heavy J Gene", "CDRH3"
        )

    def test_analysis_keeps_exact_distance_classes_separate(self):
        result = analyse(self._sample(), self._database())
        self.assertEqual([item.unique_clones for item in result.exact_summaries], [1, 1, 1])
        self.assertEqual([item.total_reads for item in result.exact_summaries], [10, 5, 3])
        self.assertEqual([item.unique_clones for item in result.cumulative_summaries], [1, 2, 3])
        self.assertEqual(result.matched_clone_count, 3)

    def test_export_contains_required_sheets(self):
        result = analyse(self._sample(), self._database())
        output = export_result_xlsx(result, test_path("result.xlsx"))
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Summary", "Method", "Input QC", "Matched Clones"])
            self.assertEqual(workbook["Matched Clones"]["A2"].value, "LV0")
            self.assertEqual(workbook["Method"]["B2"].value, ALGORITHM_VERSION)
            self.assertEqual(workbook["Method"]["B3"].value, "kobe")
        finally:
            workbook.close()

    def test_three_modes_apply_distinct_gene_candidate_rules(self):
        database_entry = DatabaseEntry("IGHV1-3", "IGHJ4", "ABCDE", {"Name": ("KnownAb",)})

        kobe_sample = SampleData(
            Path("sample.csv"),
            "CPM",
            "S1",
            (SampleClone(("IGHV1-2", "IGHV1-3"), ("IGHJ4",), "ABCDE", 10, 100.0),),
            10,
            1,
            1,
            0,
            matching_mode=MatchingMode.KOBE,
        )
        kobe_db = DatabaseData(
            Path("db.csv"),
            (database_entry,),
            1,
            1,
            0,
            ("Name",),
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            matching_mode=MatchingMode.KOBE,
        )
        self.assertEqual(analyse(kobe_sample, kobe_db).matched_clone_count, 1)

        legacy_sample = SampleData(
            Path("sample.csv"),
            "CPM",
            "S1",
            (SampleClone(("IGHV1-2 // IGHV1-3",), ("IGHJ4",), "ABCDE", 10, 100.0),),
            10,
            1,
            1,
            0,
            matching_mode=MatchingMode.LEGACY,
        )
        legacy_db = DatabaseData(
            Path("db.csv"),
            (database_entry,),
            1,
            1,
            0,
            ("Name",),
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            matching_mode=MatchingMode.LEGACY,
        )
        self.assertEqual(analyse(legacy_sample, legacy_db).matched_clone_count, 0)

        cdr3_sample = SampleData(
            Path("sample.csv"),
            "CPM",
            "S1",
            (SampleClone((), (), "ABCDE", 10, 100.0),),
            10,
            1,
            1,
            0,
            matching_mode=MatchingMode.CDR3_ONLY,
        )
        cdr3_db = DatabaseData(
            Path("db.csv"),
            (database_entry,),
            1,
            1,
            0,
            ("Name",),
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            matching_mode=MatchingMode.CDR3_ONLY,
        )
        self.assertEqual(analyse(cdr3_sample, cdr3_db).matched_clone_count, 1)

    def test_legacy_retains_all_within_two_but_uses_minimum_distance(self):
        sample = SampleData(
            Path("sample.xlsx"),
            "RG",
            "S1",
            (SampleClone(("IGHV1-2",), ("IGHJ4",), "ABCDE", 10, 100.0),),
            10,
            1,
            1,
            0,
            matching_mode=MatchingMode.LEGACY,
        )
        database = DatabaseData(
            Path("db.csv"),
            (
                DatabaseEntry("IGHV1-2", "IGHJ4", "ABCDE", {"Name": ("Exact",)}),
                DatabaseEntry("IGHV1-2", "IGHJ4", "ABCXE", {"Name": ("Near",)}),
            ),
            2,
            2,
            0,
            ("Name",),
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            matching_mode=MatchingMode.LEGACY,
        )
        result = analyse(sample, database)
        self.assertEqual(result.matches[0].distance, 0)
        self.assertEqual(result.matches[0].entry_distances, (0, 1))
        self.assertEqual(result.matches[0].combined_annotation("Name"), "Exact | Near")

    def test_cdr3_deletion_index_matches_brute_force(self):
        sequences = tuple(
            "".join(chars)
            for length in range(1, 5)
            for chars in product("ABC", repeat=length)
        )
        sample = SampleData(
            Path("sample.csv"),
            "CPM",
            "S1",
            tuple(
                SampleClone((), (), sequence, 1, 100.0 / len(sequences))
                for sequence in sequences
            ),
            len(sequences),
            len(sequences),
            len(sequences),
            0,
            matching_mode=MatchingMode.CDR3_ONLY,
        )
        database_sequences = sequences[::7]
        database = DatabaseData(
            Path("db.csv"),
            tuple(DatabaseEntry("", "", sequence) for sequence in database_sequences),
            len(database_sequences),
            len(database_sequences),
            0,
            (),
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            matching_mode=MatchingMode.CDR3_ONLY,
        )
        result = analyse(sample, database)
        indexed_distances = {match.clone.cdr3_aa: match.distance for match in result.matches}
        brute_force_distances = {}
        for sequence in sequences:
            distance = min(
                bounded_levenshtein(sequence, database_sequence, 2)
                for database_sequence in database_sequences
            )
            if distance <= 2:
                brute_force_distances[sequence] = distance
        self.assertEqual(indexed_distances, brute_force_distances)


if __name__ == "__main__":
    unittest.main()
