import csv
from hashlib import sha256
from pathlib import Path
import unittest

from openpyxl import load_workbook

from qasas.engine import analyse
from qasas.loaders import load_database, load_sample
from qasas.modes import MatchingMode
from qasas.timecourse import TimepointSpec, analyse_timecourse, parse_day, validate_timepoints
from qasas.timecourse_export import TIMECOURSE_VERSION, export_timecourse_xlsx
from tests.support import test_path


class TimeCourseTests(unittest.TestCase):
    def _database(self) -> Path:
        path = test_path("timecourse_database.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["Name", "Heavy V Gene", "Heavy J Gene", "CDRH3"],
            )
            writer.writeheader()
            writer.writerow(
                {
                    "Name": "KnownAb",
                    "Heavy V Gene": "IGHV1-2*01",
                    "Heavy J Gene": "IGHJ4*01",
                    "CDRH3": "CABCDEW",
                }
            )
        return path

    def _sample(self, name: str, exact_reads: int, lv1_reads: int) -> Path:
        path = test_path(name)
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Vseg", "Jseg", "CDR3", "Counts"])
            writer.writeheader()
            writer.writerow(
                {
                    "Vseg": "IGHV1-2*02",
                    "Jseg": "IGHJ4*02",
                    "CDR3": "CABCDEW",
                    "Counts": exact_reads,
                }
            )
            writer.writerow(
                {
                    "Vseg": "IGHV1-2*02",
                    "Jseg": "IGHJ4*02",
                    "CDR3": "CABCXEW",
                    "Counts": lv1_reads,
                }
            )
        return path

    def test_day_accepts_negative_zero_positive_and_decimal(self):
        self.assertEqual(parse_day("-7.5"), -7.5)
        self.assertEqual(parse_day("0"), 0.0)
        self.assertEqual(parse_day("14"), 14.0)
        with self.assertRaises(ValueError):
            parse_day("NaN")
        with self.assertRaises(ValueError):
            parse_day("Day14")

    def test_timepoints_are_numeric_sorted_and_duplicate_day_is_rejected(self):
        first = self._sample("tc_sort_first.csv", 4, 1)
        second = self._sample("tc_sort_second.csv", 5, 2)
        third = self._sample("tc_sort_third.csv", 6, 3)
        ordered = validate_timepoints(
            (
                TimepointSpec(14.5, "after", first),
                TimepointSpec(-3, "before", second),
                TimepointSpec(0, "baseline", third),
            )
        )
        self.assertEqual([point.day for point in ordered], [-3.0, 0.0, 14.5])
        with self.assertRaisesRegex(ValueError, "重複"):
            validate_timepoints(
                (
                    TimepointSpec(0, "a", first),
                    TimepointSpec(-0.0, "b", second),
                )
            )

    def test_each_timepoint_equals_independent_single_sample_analysis(self):
        database_path = self._database()
        day_minus = self._sample("tc_day_minus.csv", 10, 5)
        day_zero = self._sample("tc_day_zero.csv", 20, 2)
        day_decimal = self._sample("tc_day_decimal.csv", 7, 7)
        specs = (
            TimepointSpec(2.5, "decimal", day_decimal, "CPM"),
            TimepointSpec(-7, "pre", day_minus, "CPM"),
            TimepointSpec(0, "baseline", day_zero, "CPM"),
        )
        timecourse = analyse_timecourse(
            specs,
            database_path,
            MatchingMode.KOBE,
            series_name="Test series",
        )
        self.assertEqual([point.spec.day for point in timecourse.timepoints], [-7.0, 0.0, 2.5])

        database = load_database(database_path, matching_mode=MatchingMode.KOBE)
        by_path = {point.spec.sample_path: point.analysis for point in timecourse.timepoints}
        for spec in specs:
            single_sample = load_sample(spec.sample_path, spec.input_format, matching_mode=MatchingMode.KOBE)
            single_result = analyse(single_sample, database, matching_mode=MatchingMode.KOBE)
            self.assertEqual(by_path[spec.sample_path].exact_summaries, single_result.exact_summaries)
            self.assertEqual(by_path[spec.sample_path].cumulative_summaries, single_result.cumulative_summaries)
            self.assertEqual(by_path[spec.sample_path].matches, single_result.matches)

    def test_timecourse_export_contains_reproducible_longitudinal_tables(self):
        database_path = self._database()
        first = self._sample("tc_export_first.csv", 10, 5)
        second = self._sample("tc_export_second.csv", 20, 2)
        result = analyse_timecourse(
            (
                TimepointSpec(-1, "pre", first, "CPM"),
                TimepointSpec(0.5, "post", second, "CPM"),
            ),
            database_path,
            series_name="Export series",
        )
        output = export_timecourse_xlsx(result, test_path("timecourse_result.xlsx"))
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Time Course Summary", "Long Summary", "Sample QC", "Matched Clones", "Method"],
            )
            self.assertEqual(workbook["Time Course Summary"]["A6"].value, -1)
            self.assertEqual(workbook["Time Course Summary"]["A7"].value, 0.5)
            method_values = {
                workbook["Method"].cell(row=row, column=1).value:
                workbook["Method"].cell(row=row, column=2).value
                for row in range(2, workbook["Method"].max_row + 1)
            }
            self.assertEqual(method_values["Time-course module version"], TIMECOURSE_VERSION)
            self.assertEqual(method_values["Duplicate Day"], "Rejected; no automatic averaging or pooling")
            self.assertEqual(
                method_values["Common database SHA-256"],
                sha256(database_path.read_bytes()).hexdigest(),
            )
            self.assertRegex(method_values["Application source SHA-256"], r"^[0-9a-f]{64}$")
            qc = workbook["Sample QC"]
            qc_headers = [qc.cell(row=1, column=column).value for column in range(1, qc.max_column + 1)]
            sha_column = qc_headers.index("Repertoire SHA-256") + 1
            expected_sample_hashes = {
                sha256(first.read_bytes()).hexdigest(),
                sha256(second.read_bytes()).hexdigest(),
            }
            actual_sample_hashes = {
                qc.cell(row=row, column=sha_column).value for row in range(2, qc.max_row + 1)
            }
            self.assertEqual(actual_sample_hashes, expected_sample_hashes)
            self.assertEqual(workbook["Long Summary"].max_row, 13)
        finally:
            workbook.close()

    def test_timecourse_excel_charts_use_numeric_day_axis(self):
        database_path = self._database()
        day_zero = self._sample("tc_chart_day_zero.csv", 10, 5)
        day_one = self._sample("tc_chart_day_one.csv", 20, 2)
        day_fourteen = self._sample("tc_chart_day_fourteen.csv", 30, 1)
        result = analyse_timecourse(
            (
                TimepointSpec(14, "day 14", day_fourteen, "CPM"),
                TimepointSpec(0, "day 0", day_zero, "CPM"),
                TimepointSpec(1, "day 1", day_one, "CPM"),
            ),
            database_path,
            series_name="Numeric day axis",
        )
        output = export_timecourse_xlsx(result, test_path("timecourse_numeric_axis.xlsx"))
        workbook = load_workbook(output, read_only=False, data_only=False)
        try:
            summary = workbook["Time Course Summary"]
            self.assertEqual(
                [summary.cell(row=row, column=1).value for row in range(6, 9)],
                [0, 1, 14],
            )
            self.assertEqual(len(summary._charts), 2)
            for chart in summary._charts:
                self.assertEqual(chart.tagname, "scatterChart")
                self.assertEqual(chart.scatterStyle, "lineMarker")
                self.assertEqual(chart.x_axis.axPos, "b")
                self.assertEqual(chart.y_axis.axPos, "l")
                self.assertEqual(len(chart.series), 3)
                for series in chart.series:
                    self.assertEqual(
                        series.xVal.numRef.f,
                        "'Time Course Summary'!$A$6:$A$8",
                    )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()

