import csv
from pathlib import Path
import re
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook

from qasas.loaders import detect_sample_format, load_cpm, load_database, load_rg
from qasas.modes import MatchingMode
from tests.support import test_path


class LoaderTests(unittest.TestCase):
    def test_cpm_loader_aggregates_and_recalculates_frequency(self):
        path = test_path("cpm_sample.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Vseg", "Jseg", "CDR3", "Counts", "Frequency(%)"])
            writer.writeheader()
            writer.writerow({"Vseg": "IGHV3-23*01", "Jseg": "IGHJ4*02", "CDR3": "CARDRW", "Counts": 3})
            writer.writerow({"Vseg": "IGHV3-23*02", "Jseg": "IGHJ4*01", "CDR3": "CARDRW", "Counts": 2})
            writer.writerow({"Vseg": "IGHV1-2*01", "Jseg": "IGHJ6*01", "CDR3": "CQQQW", "Counts": 5})
        self.assertEqual(detect_sample_format(path), "CPM")
        sample = load_cpm(path)
        self.assertEqual(len(sample.clones), 2)
        self.assertEqual(sample.total_reads, 10)
        self.assertEqual(sample.clones[0].frequency_percent, 50.0)
        self.assertEqual(sum(clone.frequency_percent for clone in sample.clones), 100.0)

    def test_rg_loader_uses_back_data_and_in_frame_only(self):
        path = test_path("rg_sample.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Back_data"
        sheet["A1"] = "Sample ID"
        sheet["C1"] = "RG001"
        sheet["A2"] = "In-frame reads"
        sheet["C2"] = 12
        sheet.cell(10, 7, "IGHV3-23*01")
        sheet.cell(10, 11, "IGHJ4*02")
        sheet.cell(10, 15, "CARDRW")
        sheet.cell(10, 16, "in-frame")
        sheet.cell(10, 17, 7)
        sheet.cell(11, 7, "IGHV3-23*02")
        sheet.cell(11, 11, "IGHJ4*01")
        sheet.cell(11, 15, "CARDRW")
        sheet.cell(11, 16, "in-frame")
        sheet.cell(11, 17, 5)
        sheet.cell(12, 7, "IGHV1-2*01")
        sheet.cell(12, 11, "IGHJ6*01")
        sheet.cell(12, 15, "COUTW")
        sheet.cell(12, 16, "out-of-frame")
        sheet.cell(12, 17, 99)
        workbook.save(path)
        sample = load_rg(path)
        self.assertEqual(sample.sample_id, "RG001")
        self.assertEqual(len(sample.clones), 1)
        self.assertEqual(sample.total_reads, 12)
        self.assertEqual(sample.listed_reads, 12)
        self.assertEqual(sample.source_rows, 3)
        self.assertEqual(sample.skipped_rows, 1)

    def test_rg_loader_does_not_trust_truncated_worksheet_dimension(self):
        path = test_path("rg_truncated_dimension.xlsx")
        repacked = test_path("rg_truncated_dimension_repacked.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Back_data"
        sheet["A1"] = "3:Sample ID"
        sheet["C1"] = "RG-DIMENSION"
        sheet["B2"] = "In frame"
        sheet["C2"] = 12
        sheet.cell(10, 7, "IGHV3-23*01")
        sheet.cell(10, 11, "IGHJ4*02")
        sheet.cell(10, 15, "CARDRW")
        sheet.cell(10, 16, "in-frame")
        sheet.cell(10, 17, 7)
        sheet.cell(300, 7, "IGHV1-2*01")
        sheet.cell(300, 11, "IGHJ6*01")
        sheet.cell(300, 15, "CQQQW")
        sheet.cell(300, 16, "in-frame")
        sheet.cell(300, 17, 5)
        workbook.save(path)

        with ZipFile(path, "r") as source, ZipFile(repacked, "w", ZIP_DEFLATED) as destination:
            for member in source.infolist():
                payload = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    payload = re.sub(
                        br'<dimension ref="[^"]+"',
                        b'<dimension ref="A1:Q10"',
                        payload,
                        count=1,
                    )
                destination.writestr(member, payload)
        repacked.replace(path)

        read_only_workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            self.assertEqual(read_only_workbook["Back_data"].max_row, 10)
        finally:
            read_only_workbook.close()

        sample = load_rg(path)
        self.assertEqual(sample.source_rows, 2)
        self.assertEqual(len(sample.clones), 2)
        self.assertEqual(sample.listed_reads, 12)
        self.assertEqual(sample.total_reads, 12)

    def test_legacy_rg_frequency_uses_listed_reads(self):
        path = test_path("rg_denominator.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Back_data"
        sheet["A1"] = "Sample ID"
        sheet["C1"] = "RG-DENOMINATOR"
        sheet["A2"] = "In-frame reads"
        sheet["C2"] = 100
        sheet.cell(10, 7, "IGHV3-23*01")
        sheet.cell(10, 11, "IGHJ4*02")
        sheet.cell(10, 15, "CARDRW")
        sheet.cell(10, 16, "in-frame")
        sheet.cell(10, 17, 95)
        workbook.save(path)

        legacy = load_rg(path, matching_mode=MatchingMode.LEGACY)
        kobe = load_rg(path, matching_mode=MatchingMode.KOBE)
        self.assertEqual(legacy.listed_reads, 95)
        self.assertEqual(legacy.total_reads, 95)
        self.assertEqual(kobe.listed_reads, 95)
        self.assertEqual(kobe.total_reads, 100)

    def test_database_loader_deduplicates_v_j_cdr3_and_merges_annotations(self):
        path = test_path("database.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["Name", "Heavy V Gene", "Heavy J Gene", "CDRH3", "Binds to"],
            )
            writer.writeheader()
            writer.writerow(
                {"Name": "Ab1", "Heavy V Gene": "IGHV3-23*01 (Human)", "Heavy J Gene": "IGHJ4*02", "CDRH3": "ARDR", "Binds to": "A"}
            )
            writer.writerow(
                {"Name": "Ab2", "Heavy V Gene": "IGHV3-23*02", "Heavy J Gene": "IGHJ4*01", "CDRH3": "ARDR", "Binds to": "B"}
            )
        database = load_database(path)
        self.assertEqual(database.usable_rows, 2)
        self.assertEqual(len(database.entries), 1)
        self.assertEqual(database.entries[0].annotations["Name"], ("Ab1", "Ab2"))

    def test_cdr3_only_aggregates_same_cdr3_across_vj_calls(self):
        path = test_path("cdr3_only_sample.csv")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Vseg", "Jseg", "CDR3", "Counts"])
            writer.writeheader()
            writer.writerow({"Vseg": "IGHV1-2*01", "Jseg": "IGHJ4*01", "CDR3": "CARDRW", "Counts": 3})
            writer.writerow({"Vseg": "IGHV3-23*01", "Jseg": "IGHJ6*01", "CDR3": "CARDRW", "Counts": 7})
        sample = load_cpm(path, matching_mode=MatchingMode.CDR3_ONLY)
        self.assertEqual(len(sample.clones), 1)
        self.assertEqual(sample.clones[0].reads, 10)
        self.assertEqual(sample.clones[0].cdr3_aa, "ARDR")
        self.assertEqual(sample.clones[0].v_genes, ("IGHV1-2", "IGHV3-23"))

    def test_legacy_loader_preserves_gene_text_and_applies_cov_abdab_rules(self):
        sample_path = test_path("legacy_sample.csv")
        with sample_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Vseg", "Jseg", "CDR3", "Counts"])
            writer.writeheader()
            writer.writerow(
                {
                    "Vseg": "IGHV1-2*01 // IGHV1-3*01",
                    "Jseg": "IGHJ4*02",
                    "CDR3": "CABCDEW",
                    "Counts": 5,
                }
            )
        sample = load_cpm(sample_path, matching_mode=MatchingMode.LEGACY)
        self.assertEqual(sample.clones[0].v_genes, ("IGHV1-2*01 // IGHV1-3*01",))
        self.assertEqual(sample.clones[0].j_genes, ("IGHJ4*02",))
        self.assertEqual(sample.clones[0].cdr3_aa, "CABCDEW")

        database_path = test_path("legacy_database.csv")
        fieldnames = [
            "Name",
            "Heavy V Gene",
            "Heavy J Gene",
            "CDRH3",
            "Binds to",
            "Doesn't Bind to",
            "Neutralising Vs",
            "Not Neutralising Vs",
        ]
        with database_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerow(
                {
                    "Name": "HumanBound",
                    "Heavy V Gene": "IGHV1-2*01 (Human)",
                    "Heavy J Gene": "IGHJ4*02 (Human)",
                    "CDRH3": "ABCDE",
                    "Binds to": "WT",
                }
            )
            writer.writerow(
                {
                    "Name": "MouseBound",
                    "Heavy V Gene": "IGHV1-2*01 (Mouse)",
                    "Heavy J Gene": "IGHJ4*02 (Mouse)",
                    "CDRH3": "ABCDE",
                    "Binds to": "WT",
                }
            )
            writer.writerow(
                {
                    "Name": "HumanUnannotated",
                    "Heavy V Gene": "IGHV1-2*01 (Human)",
                    "Heavy J Gene": "IGHJ4*02 (Human)",
                    "CDRH3": "ABCDE",
                }
            )
        database = load_database(database_path, matching_mode=MatchingMode.LEGACY)
        self.assertEqual(database.usable_rows, 1)
        self.assertEqual(len(database.entries), 1)
        self.assertEqual(database.entries[0].v_gene, "IGHV1-2*01")
        self.assertEqual(database.entries[0].j_gene, "IGHJ4*02")
        self.assertEqual(database.entries[0].cdr3_aa, "CABCDEW")


if __name__ == "__main__":
    unittest.main()
