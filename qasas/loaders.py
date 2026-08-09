from __future__ import annotations

import csv
import re
from collections import OrderedDict
from collections.abc import Callable
from itertools import product
from pathlib import Path

from openpyxl import load_workbook

from .models import DatabaseData, DatabaseEntry, SampleClone, SampleData
from .modes import MatchingMode, mode_specification, parse_matching_mode
from .normalize import (
    clean_text,
    normalize_cdr3,
    normalize_gene_options,
    parse_positive_int,
)


StatusCallback = Callable[[str], None] | None


def _status(callback: StatusCallback, message: str) -> None:
    if callback:
        callback(message)


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp932", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.DictReader(handle, dialect=dialect)
                fieldnames = [str(name).strip() for name in (reader.fieldnames or [])]
                rows = []
                for row in reader:
                    rows.append(
                        {
                            str(key).strip(): "" if value is None else str(value).strip()
                            for key, value in row.items()
                            if key is not None
                        }
                    )
                return fieldnames, rows
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error:
        raise ValueError(f"CSVの文字コードを判定できません: {path}") from last_error
    raise ValueError(f"CSVを読み込めません: {path}")


def _header_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def _find_column(headers: list[str], aliases: tuple[str, ...], label: str) -> str:
    normalized = {_header_key(header): header for header in headers}
    for alias in aliases:
        hit = normalized.get(_header_key(alias))
        if hit:
            return hit
    raise ValueError(f"{label}列が見つかりません。列名: {', '.join(headers)}")


def _find_optional_column(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    normalized = {_header_key(header): header for header in headers}
    for alias in aliases:
        hit = normalized.get(_header_key(alias))
        if hit:
            return hit
    return None


def _legacy_gene_call(value: object, *, database_side: bool = False) -> str:
    """Keep an old-QASAS gene call intact except for its exact species suffix."""

    text = clean_text(value)
    if database_side and text.endswith(" (Human)"):
        text = text[: -len(" (Human)")]
    return text


def _legacy_annotation_set(value: object) -> set[str]:
    output: set[str] = set()
    for item in clean_text(value).split(";"):
        normalized = item.replace("(weak)", "").replace(" ", "")
        if normalized:
            output.add(normalized)
    return output


def _legacy_is_antigen_specific(
    row: dict[str, str],
    bind_column: str | None,
    not_bind_column: str | None,
    neut_column: str | None,
    not_neut_column: str | None,
) -> bool:
    """Reproduce the old Bind-or-Neut row filter when those columns exist."""

    columns = (bind_column, not_bind_column, neut_column, not_neut_column)
    if not any(columns):
        return True
    bind = _legacy_annotation_set(row.get(bind_column, "") if bind_column else "")
    not_bind = _legacy_annotation_set(row.get(not_bind_column, "") if not_bind_column else "")
    neut = _legacy_annotation_set(row.get(neut_column, "") if neut_column else "")
    not_neut = _legacy_annotation_set(row.get(not_neut_column, "") if not_neut_column else "")
    bind_var = bind - not_bind
    neut_var = neut - not_neut
    return bool(bind_var | neut_var)


def detect_sample_format(path: str | Path) -> str:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return "RG"
    if suffix in {".csv", ".tsv", ".txt"}:
        headers, _ = _read_csv(source)
        keys = {_header_key(header) for header in headers}
        if {_header_key("Vseg"), _header_key("Jseg"), _header_key("CDR3")} <= keys:
            return "CPM"
    raise ValueError("入力形式を自動判定できません。CPM様式またはRG様式を指定してください。")


def _build_sample_data(
    *,
    path: Path,
    input_format: str,
    sample_id: str,
    raw_clones: list[tuple[object, object, object, int]],
    source_rows: int,
    skipped_rows: int,
    matching_mode: MatchingMode | str = MatchingMode.KOBE,
    metadata: dict[str, str] | None = None,
    total_reads_override: int | None = None,
) -> SampleData:
    mode = parse_matching_mode(matching_mode)
    spec = mode_specification(mode)
    aggregate: OrderedDict[
        tuple[tuple[str, ...], tuple[str, ...], str],
        dict[str, object],
    ] = OrderedDict()

    accepted_rows = 0
    for raw_v, raw_j, raw_cdr3, reads in raw_clones:
        if mode is MatchingMode.LEGACY:
            v_text = _legacy_gene_call(raw_v)
            j_text = _legacy_gene_call(raw_j)
            v_options = (v_text,) if v_text else ()
            j_options = (j_text,) if j_text else ()
            cdr3 = clean_text(raw_cdr3)
            required_genes_present = bool(v_options and j_options)
        else:
            v_options = normalize_gene_options(raw_v, "IGHV")
            j_options = normalize_gene_options(raw_j, "IGHJ")
            cdr3 = normalize_cdr3(raw_cdr3)
            required_genes_present = mode is MatchingMode.CDR3_ONLY or bool(v_options and j_options)
        if not required_genes_present or not cdr3 or reads <= 0:
            skipped_rows += 1
            continue
        accepted_rows += 1
        key = ((), (), cdr3) if mode is MatchingMode.CDR3_ONLY else (v_options, j_options, cdr3)
        if key not in aggregate:
            aggregate[key] = {
                "reads": 0,
                "rows": 0,
                "v_options": OrderedDict(),
                "j_options": OrderedDict(),
                "raw_v": OrderedDict(),
                "raw_j": OrderedDict(),
                "raw_cdr3": OrderedDict(),
            }
        item = aggregate[key]
        item["reads"] = int(item["reads"]) + reads
        item["rows"] = int(item["rows"]) + 1
        for gene in v_options:
            item["v_options"][gene] = None
        for gene in j_options:
            item["j_options"][gene] = None
        for bucket, raw_value in (
            (item["raw_v"], clean_text(raw_v)),
            (item["raw_j"], clean_text(raw_j)),
            (item["raw_cdr3"], clean_text(raw_cdr3)),
        ):
            if raw_value:
                bucket[raw_value] = None

    listed_reads = sum(int(item["reads"]) for item in aggregate.values())
    total_reads = (
        total_reads_override
        if total_reads_override is not None and total_reads_override >= listed_reads
        else listed_reads
    )
    clones: list[SampleClone] = []
    for (key_v_options, key_j_options, cdr3), item in aggregate.items():
        reads = int(item["reads"])
        frequency = (reads / total_reads * 100.0) if total_reads else 0.0
        if mode is MatchingMode.CDR3_ONLY:
            v_options = tuple(item["v_options"].keys())
            j_options = tuple(item["j_options"].keys())
        else:
            v_options = key_v_options
            j_options = key_j_options
        clones.append(
            SampleClone(
                v_genes=v_options,
                j_genes=j_options,
                cdr3_aa=cdr3,
                reads=reads,
                frequency_percent=frequency,
                raw_v_gene=" | ".join(item["raw_v"].keys()),
                raw_j_gene=" | ".join(item["raw_j"].keys()),
                raw_cdr3_aa=" | ".join(item["raw_cdr3"].keys()),
                source_rows=int(item["rows"]),
            )
        )
    clones.sort(key=lambda clone: (-clone.reads, clone.display_v_gene, clone.cdr3_aa))
    output_metadata = dict(metadata or {})
    output_metadata.update(
        {
            "Matching mode code": mode.value,
            "Matching mode": spec.label,
            "Sample clone key": spec.sample_clone_key,
            "Sample gene handling": spec.sample_gene_handling,
            "CDR3 handling": spec.cdr3_handling,
        }
    )
    return SampleData(
        source_path=path,
        input_format=input_format,
        sample_id=sample_id,
        clones=tuple(clones),
        total_reads=total_reads,
        source_rows=source_rows,
        accepted_rows=accepted_rows,
        skipped_rows=skipped_rows,
        metadata=output_metadata,
        matching_mode=mode,
    )


def load_cpm(
    path: str | Path,
    status_callback: StatusCallback = None,
    matching_mode: MatchingMode | str = MatchingMode.KOBE,
) -> SampleData:
    source = Path(path)
    _status(status_callback, "CPM様式の検体データを読み込んでいます…")
    headers, rows = _read_csv(source)
    v_column = _find_column(headers, ("Vseg", "V gene", "IGHV"), "IGHV")
    j_column = _find_column(headers, ("Jseg", "J gene", "IGHJ"), "IGHJ")
    cdr3_column = _find_column(headers, ("CDR3", "CDR3 AA", "junction_aa"), "CDR3")
    count_column = _find_column(
        headers,
        ("Counts", "Count", "Reads", "Read count", "DUPCOUNT"),
        "リード数",
    )
    raw_clones = [
        (row.get(v_column), row.get(j_column), row.get(cdr3_column), parse_positive_int(row.get(count_column)))
        for row in rows
    ]
    return _build_sample_data(
        path=source,
        input_format="CPM",
        sample_id=source.stem,
        raw_clones=raw_clones,
        source_rows=len(rows),
        skipped_rows=0,
        matching_mode=matching_mode,
        metadata={"Count column": count_column},
    )


def _normalize_label(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean_text(value).lower())


def _extract_rg_metadata(rows: list[tuple[object, ...]]) -> dict[str, str]:
    metadata: dict[str, str] = {}
    known_labels = {
        "sampleid": "Sample ID",
        "totalreads": "Total reads",
        "assignedreads": "Assigned reads",
        "inframereads": "In-frame reads",
        "uniqueinframereads": "Unique in-frame reads",
        "uniquereadsinframe": "Unique in-frame reads",
        "inframe": "In-frame reads",
    }
    for row in rows[:25]:
        for index, value in enumerate(row[:5]):
            normalized = re.sub(r"^\d+", "", _normalize_label(value))
            if normalized not in known_labels:
                continue
            following = [clean_text(cell) for cell in row[index + 1 :] if clean_text(cell)]
            if following:
                metadata[known_labels[normalized]] = following[0]
    return metadata


def load_rg(
    path: str | Path,
    status_callback: StatusCallback = None,
    matching_mode: MatchingMode | str = MatchingMode.KOBE,
) -> SampleData:
    source = Path(path)
    mode = parse_matching_mode(matching_mode)
    _status(status_callback, "RG様式のBack_data全件領域を読み込んでいます…")
    # RG reports can contain tens of thousands of Back_data rows while their
    # cached worksheet dimension says only A1:BD212.  openpyxl's read-only
    # mode trusts that stale dimension and silently truncates the repertoire.
    # Normal mode determines the range from the actual cells.
    workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        sheet_name = next(
            (name for name in workbook.sheetnames if _header_key(name) == _header_key("Back_data")),
            None,
        )
        if not sheet_name:
            raise ValueError("RG様式に必要なBack_dataシートが見つかりません。")
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(min_col=1, max_col=17, values_only=True))
    finally:
        workbook.close()

    metadata = _extract_rg_metadata(rows)
    raw_clones: list[tuple[object, object, object, int]] = []
    out_of_frame_rows = 0
    non_data_rows = 0
    candidate_rows = 0
    for row in rows:
        frame = _normalize_label(row[15] if len(row) > 15 else None)
        if frame not in {"inframe", "outofframe"}:
            continue
        candidate_rows += 1
        if frame != "inframe":
            out_of_frame_rows += 1
            continue
        reads = parse_positive_int(row[16] if len(row) > 16 else None)
        if reads <= 0:
            non_data_rows += 1
            continue
        raw_clones.append((row[6], row[10], row[14], reads))

    if not candidate_rows:
        raise ValueError(
            "Back_dataの全件領域を認識できません。G=IGHV、K=IGHJ、O=CDR3、P=frame、Q=Readsを確認してください。"
        )
    sample_id = metadata.get("Sample ID", source.stem)
    metadata = dict(metadata)
    metadata["RG data region"] = "Back_data 7:All Data"
    metadata["Out-of-frame rows excluded"] = str(out_of_frame_rows)
    listed_reads = sum(item[3] for item in raw_clones)
    in_frame_denominator = parse_positive_int(metadata.get("In-frame reads"))
    if in_frame_denominator > 0 and listed_reads < in_frame_denominator * 0.95:
        raise ValueError(
            "RGのBack_dataが途中までしか読み込まれていません。"
            f"列挙リード {listed_reads:,} / In frame総リード {in_frame_denominator:,}。"
            "元のRG Excelが全件を含むか確認してください。"
        )
    if mode is MatchingMode.LEGACY:
        in_frame_denominator = listed_reads
        metadata["Frequency denominator"] = f"Legacy listed in-frame reads: {listed_reads}"
    elif in_frame_denominator >= listed_reads and in_frame_denominator > 0:
        metadata["Frequency denominator"] = f"In-frame reads: {in_frame_denominator}"
    else:
        in_frame_denominator = listed_reads
        metadata["Frequency denominator"] = f"Listed in-frame reads: {listed_reads}"
    metadata["Reads represented by listed clones"] = str(listed_reads)
    metadata["Listed in-frame rows"] = str(len(raw_clones))
    return _build_sample_data(
        path=source,
        input_format="RG",
        sample_id=sample_id,
        raw_clones=raw_clones,
        source_rows=candidate_rows,
        skipped_rows=out_of_frame_rows + non_data_rows,
        matching_mode=mode,
        metadata=metadata,
        total_reads_override=in_frame_denominator,
    )


def load_sample(
    path: str | Path,
    input_format: str = "AUTO",
    status_callback: StatusCallback = None,
    matching_mode: MatchingMode | str = MatchingMode.KOBE,
) -> SampleData:
    requested = input_format.strip().upper()
    if requested in {"AUTO", "自動判定"}:
        requested = detect_sample_format(path)
    if requested in {"CPM", "CPM様式"}:
        return load_cpm(path, status_callback, matching_mode)
    if requested in {"RG", "RG様式"}:
        return load_rg(path, status_callback, matching_mode)
    raise ValueError(f"未対応の入力形式です: {input_format}")


def load_database(
    path: str | Path,
    status_callback: StatusCallback = None,
    matching_mode: MatchingMode | str = MatchingMode.KOBE,
) -> DatabaseData:
    source = Path(path)
    mode = parse_matching_mode(matching_mode)
    spec = mode_specification(mode)
    _status(status_callback, "抗原結合性データベースを読み込んでいます…")
    headers, rows = _read_csv(source)
    v_column = _find_column(
        headers,
        ("Heavy V Gene", "IGHV", "V gene", "Vseg"),
        "データベースIGHV",
    )
    j_column = _find_column(
        headers,
        ("Heavy J Gene", "IGHJ", "J gene", "Jseg"),
        "データベースIGHJ",
    )
    cdr3_column = _find_column(
        headers,
        ("CDRH3", "CDR3", "CDR3 AA", "junction_aa"),
        "データベースCDR3",
    )
    source_annotation_columns = tuple(
        header for header in headers if header not in {v_column, j_column, cdr3_column}
    )
    annotation_columns = (
        ("DB IGHV source", "DB IGHJ source") + source_annotation_columns
        if mode is MatchingMode.CDR3_ONLY
        else source_annotation_columns
    )

    raw_cdr3_values = [clean_text(row.get(cdr3_column)).upper() for row in rows]
    nonempty_cdr3_values = [value for value in raw_cdr3_values if value]
    anchored_count = sum(
        len(value) > 2 and value.startswith("C") and value.endswith("W")
        for value in nonempty_cdr3_values
    )
    # Kobe/CDR3-only harmonize anchors. Legacy instead adds C and W unconditionally.
    trim_database_anchors = bool(nonempty_cdr3_values) and anchored_count / len(nonempty_cdr3_values) >= 0.8

    human_tagged_database = any(
        "Human" in clean_text(row.get(v_column)) or "Human" in clean_text(row.get(j_column))
        for row in rows
    )
    bind_column = _find_optional_column(headers, ("Binds to", "Bind"))
    not_bind_column = _find_optional_column(headers, ("Doesn't Bind to", "Does not Bind to", "notBind"))
    neut_column = _find_optional_column(headers, ("Neutralising Vs", "Neutralizing Vs", "Neut"))
    not_neut_column = _find_optional_column(
        headers,
        ("Not Neutralising Vs", "Not Neutralizing Vs", "notNeut"),
    )
    has_legacy_antigen_columns = any((bind_column, not_bind_column, neut_column, not_neut_column))

    aggregates: OrderedDict[tuple[str, str, str], dict[str, object]] = OrderedDict()
    usable_rows = 0
    skipped_rows = 0
    legacy_nonhuman_rows = 0
    legacy_unannotated_rows = 0
    for row in rows:
        raw_v = clean_text(row.get(v_column))
        raw_j = clean_text(row.get(j_column))
        raw_cdr3 = clean_text(row.get(cdr3_column))
        if mode is MatchingMode.LEGACY:
            if human_tagged_database and not ("Human" in raw_v and "Human" in raw_j):
                skipped_rows += 1
                legacy_nonhuman_rows += 1
                continue
            if not _legacy_is_antigen_specific(
                row,
                bind_column,
                not_bind_column,
                neut_column,
                not_neut_column,
            ):
                skipped_rows += 1
                legacy_unannotated_rows += 1
                continue
            v_gene = _legacy_gene_call(raw_v, database_side=True)
            j_gene = _legacy_gene_call(raw_j, database_side=True)
            cdr3 = f"C{raw_cdr3}W" if raw_cdr3 else ""
            keys = ((v_gene, j_gene, cdr3),) if v_gene and j_gene and cdr3 else ()
        elif mode is MatchingMode.CDR3_ONLY:
            cdr3 = normalize_cdr3(raw_cdr3, trim_anchors=trim_database_anchors)
            keys = (("", "", cdr3),) if cdr3 else ()
        else:
            v_options = normalize_gene_options(raw_v, "IGHV")
            j_options = normalize_gene_options(raw_j, "IGHJ")
            cdr3 = normalize_cdr3(raw_cdr3, trim_anchors=trim_database_anchors)
            keys = tuple(product(v_options, j_options))
            keys = tuple((v_gene, j_gene, cdr3) for v_gene, j_gene in keys) if cdr3 else ()

        if not keys:
            skipped_rows += 1
            continue
        usable_rows += 1
        for key in keys:
            if key not in aggregates:
                aggregates[key] = {
                    "source_rows": 0,
                    "annotations": {column: OrderedDict() for column in annotation_columns},
                }
            item = aggregates[key]
            item["source_rows"] = int(item["source_rows"]) + 1
            annotations = item["annotations"]
            assert isinstance(annotations, dict)
            for column in source_annotation_columns:
                value = clean_text(row.get(column))
                if value:
                    annotations[column][value] = None
            if mode is MatchingMode.CDR3_ONLY:
                if raw_v:
                    annotations["DB IGHV source"][raw_v] = None
                if raw_j:
                    annotations["DB IGHJ source"][raw_j] = None

    entries: list[DatabaseEntry] = []
    for (v_gene, j_gene, cdr3), item in aggregates.items():
        annotations = item["annotations"]
        assert isinstance(annotations, dict)
        entries.append(
            DatabaseEntry(
                v_gene=v_gene,
                j_gene=j_gene,
                cdr3_aa=cdr3,
                annotations={column: tuple(values.keys()) for column, values in annotations.items()},
                source_rows=int(item["source_rows"]),
            )
        )
    metadata = {
        "Matching mode code": mode.value,
        "Matching mode": spec.label,
        "Database key": spec.database_key,
        "Candidate rule": spec.candidate_rule,
        "Candidate retention": spec.candidate_retention,
        "Database CDR3 anchor convention": (
            "Legacy: add C and W unconditionally"
            if mode is MatchingMode.LEGACY
            else f"Trim anchors={trim_database_anchors} ({anchored_count}/{len(nonempty_cdr3_values)} anchored nonempty rows)"
        ),
    }
    if mode is MatchingMode.LEGACY:
        metadata.update(
            {
                "Legacy human-tag filter active": str(human_tagged_database),
                "Legacy non-human rows excluded": str(legacy_nonhuman_rows),
                "Legacy Bind-or-Neut filter active": str(has_legacy_antigen_columns),
                "Legacy rows without Bind-or-Neut excluded": str(legacy_unannotated_rows),
            }
        )

    return DatabaseData(
        source_path=source,
        entries=tuple(entries),
        source_rows=len(rows),
        usable_rows=usable_rows,
        skipped_rows=skipped_rows,
        annotation_columns=annotation_columns,
        v_column=v_column,
        j_column=j_column,
        cdr3_column=cdr3_column,
        matching_mode=mode,
        metadata=metadata,
    )
