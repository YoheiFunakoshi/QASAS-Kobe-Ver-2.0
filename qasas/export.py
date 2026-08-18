from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AnalysisResult
from .modes import ALGORITHM_VERSION, mode_specification
from .provenance import file_provenance, file_provenance_rows, runtime_provenance_rows


_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_WHITE_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int, start: int, end: int) -> None:
    for column in range(start, end + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(sheet, minimum: int = 10, maximum: int = 48) -> None:
    for column_cells in sheet.columns:
        length = 0
        for cell in column_cells:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            maximum, max(minimum, length + 2)
        )


def _annotation_order(result: AnalysisResult) -> list[str]:
    preferred = [
        "Name",
        "Binds to",
        "Protein + Epitope",
        "Neutralising Vs",
        "Not Neutralising Vs",
        "Sources",
    ]
    available = list(result.database.annotation_columns)
    ordered = [column for column in preferred if column in available]
    ordered.extend(column for column in available if column not in ordered)
    return ordered


def export_result_xlsx(result: AnalysisResult, path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    spec = mode_specification(result.matching_mode)
    sample_provenance = file_provenance(result.sample.source_path)
    database_provenance = file_provenance(result.database.source_path)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    summary["A1"] = "QASAS Kobe Ver 2.0"
    summary["A1"].font = Font(size=16, bold=True, color="17365D")
    info_rows = [
        ("Analysis date", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Algorithm version", ALGORITHM_VERSION),
        ("Matching mode code", result.matching_mode.value),
        ("Matching mode", spec.label),
        ("Sample ID", result.sample.sample_id),
        ("Input format", result.sample.input_format),
        ("Repertoire file", str(result.sample.source_path)),
        ("Repertoire SHA-256", sample_provenance.sha256),
        ("Database input format", result.database.metadata.get("Database input format", "Unspecified")),
        ("Database file", str(result.database.source_path)),
        ("Database SHA-256", database_provenance.sha256),
        ("Sample unique clones", len(result.sample.clones)),
        ("Sample listed reads", result.sample.listed_reads),
        ("Sample frequency denominator reads", result.sample.total_reads),
        ("Database source rows", result.database.source_rows),
        ("Database usable rows", result.database.usable_rows),
        ("Database unique matching keys", len(result.database.entries)),
    ]
    for row_number, (label, value) in enumerate(info_rows, start=3):
        summary.cell(row=row_number, column=1, value=label).font = Font(bold=True)
        summary.cell(row=row_number, column=2, value=value)

    start_row = len(info_rows) + 5
    summary.cell(start_row, 1, "Exact distance class")
    summary.cell(start_row, 2, "Unique clones")
    summary.cell(start_row, 3, "Total reads")
    summary.cell(start_row, 4, "Frequency (%)")
    _style_header(summary, start_row, 1, 4)
    for offset, item in enumerate(result.exact_summaries, start=1):
        row = start_row + offset
        summary.cell(row, 1, item.label)
        summary.cell(row, 2, item.unique_clones)
        summary.cell(row, 3, item.total_reads)
        summary.cell(row, 4, item.frequency_percent)
        summary.cell(row, 4).number_format = "0.000000000"

    cumulative_start = start_row + len(result.exact_summaries) + 3
    summary.cell(cumulative_start, 1, "Cumulative distance")
    summary.cell(cumulative_start, 2, "Unique clones")
    summary.cell(cumulative_start, 3, "Total reads")
    summary.cell(cumulative_start, 4, "Frequency (%)")
    _style_header(summary, cumulative_start, 1, 4)
    for offset, item in enumerate(result.cumulative_summaries, start=1):
        row = cumulative_start + offset
        summary.cell(row, 1, item.label)
        summary.cell(row, 2, item.unique_clones)
        summary.cell(row, 3, item.total_reads)
        summary.cell(row, 4, item.frequency_percent)
        summary.cell(row, 4).number_format = "0.000000000"

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "QASAS exact Levenshtein classes"
    chart.y_axis.title = "Unique clones"
    chart.x_axis.title = "Distance class"
    data = Reference(
        summary,
        min_col=2,
        min_row=start_row,
        max_row=start_row + len(result.exact_summaries),
    )
    categories = Reference(
        summary,
        min_col=1,
        min_row=start_row + 1,
        max_row=start_row + len(result.exact_summaries),
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    summary.add_chart(chart, "F3")

    method = workbook.create_sheet("Method")
    method.append(["Item", "Fixed rule for this result"])
    _style_header(method, 1, 1, 2)
    method_rows = [
        ("Algorithm version", ALGORITHM_VERSION),
        ("Matching mode code", result.matching_mode.value),
        ("Matching mode", spec.label),
        ("Purpose", spec.description),
        ("Sample clone key", spec.sample_clone_key),
        ("Sample gene handling", spec.sample_gene_handling),
        ("Database key", spec.database_key),
        ("Candidate selection", spec.candidate_rule),
        ("CDR3 handling", spec.cdr3_handling),
        ("Database candidate retention", spec.candidate_retention),
        ("Frequency denominator", spec.frequency_rule),
        ("Distance function", "Levenshtein distance; insertion, deletion, substitution cost = 1"),
        ("Distance classes", "Exact LV0/LV1/LV2 and cumulative ≤LV0/≤LV1/≤LV2"),
        ("Maximum distance", len(result.exact_summaries) - 1),
        ("Reproducibility warning", "Results from different matching modes must not be pooled as the same method."),
    ]
    method_rows.extend(runtime_provenance_rows())
    for row in method_rows:
        method.append(row)
    method.freeze_panes = "A2"

    qc = workbook.create_sheet("Input QC")
    qc.append(["Item", "Value"])
    _style_header(qc, 1, 1, 2)
    qc_rows = [
        *file_provenance_rows("Repertoire", sample_provenance),
        *file_provenance_rows("Database", database_provenance),
        ("Sample source rows", result.sample.source_rows),
        ("Sample accepted rows", result.sample.accepted_rows),
        ("Sample skipped rows", result.sample.skipped_rows),
        ("Sample normalized unique clones", len(result.sample.clones)),
        ("Sample reads represented by listed clones", result.sample.listed_reads),
        ("Sample frequency denominator reads", result.sample.total_reads),
        ("Database source rows", result.database.source_rows),
        ("Database usable rows", result.database.usable_rows),
        ("Database skipped rows", result.database.skipped_rows),
        ("Database normalized unique keys", len(result.database.entries)),
        ("Database V column", result.database.v_column),
        ("Database J column", result.database.j_column),
        ("Database CDR3 column", result.database.cdr3_column),
    ]
    qc_rows.extend((f"RG metadata: {key}", value) for key, value in result.sample.metadata.items())
    qc_rows.extend((f"Database metadata: {key}", value) for key, value in result.database.metadata.items())
    for row in qc_rows:
        qc.append(row)

    matched = workbook.create_sheet("Matched Clones")
    annotations = _annotation_order(result)
    headers = [
        "Distance",
        "Matching mode",
        "IGHV normalized",
        "IGHJ normalized",
        "CDR3 normalized",
        "IGHV source",
        "IGHJ source",
        "CDR3 source",
        "Reads",
        "Frequency (%)",
        "Source rows aggregated",
        "Database entries retained",
        "Database distances retained",
    ] + annotations
    matched.append(headers)
    _style_header(matched, 1, 1, len(headers))
    for match in result.matches:
        row = [
            f"LV{match.distance}",
            spec.label,
            match.clone.display_v_gene,
            match.clone.display_j_gene,
            match.clone.cdr3_aa,
            match.clone.raw_v_gene,
            match.clone.raw_j_gene,
            match.clone.raw_cdr3_aa,
            match.clone.reads,
            match.clone.frequency_percent,
            match.clone.source_rows,
            len(match.database_entries),
            " | ".join(f"LV{distance}" for distance in match.entry_distances),
        ]
        row.extend(match.combined_annotation(column) for column in annotations)
        matched.append(row)
        matched.cell(matched.max_row, 10).number_format = "0.000000000"

    matched.freeze_panes = "A2"
    matched.auto_filter.ref = matched.dimensions
    summary.freeze_panes = "A3"
    _autosize(method)
    qc.freeze_panes = "A2"
    _autosize(summary)
    _autosize(qc)
    _autosize(matched)
    workbook.save(destination)
    return destination
