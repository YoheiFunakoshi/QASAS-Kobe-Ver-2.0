from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modes import ALGORITHM_VERSION, mode_specification
from .provenance import file_provenance, runtime_provenance_rows
from .timecourse import TimeCourseResult, format_day


APP_VERSION = "QASAS Kobe Ver 2.0"
TIMECOURSE_VERSION = "QASAS-Kobe-timecourse-2.0"
_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_WHITE_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int, start: int, end: int) -> None:
    for column in range(start, end + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(sheet, minimum: int = 10, maximum: int = 48) -> None:
    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            maximum, max(minimum, length + 2)
        )


def _annotation_order(result: TimeCourseResult) -> list[str]:
    preferred = [
        "Name",
        "Binds to",
        "Protein + Epitope",
        "Neutralising Vs",
        "Not Neutralising Vs",
        "Sources",
    ]
    available = list(result.database.annotation_columns)
    ordered = [column for column in preferred if column in available]
    ordered.extend(column for column in available if column not in ordered)
    return ordered


def _timecourse_scatter_chart(
    summary,
    *,
    header_row: int,
    final_row: int,
    min_col: int,
    max_col: int,
    title: str,
    y_axis_title: str,
) -> ScatterChart:
    """Create a line-and-marker chart with numeric Day spacing on the x-axis."""

    chart = ScatterChart()
    chart.scatterStyle = "lineMarker"
    chart.title = title
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = "Day"
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    day_values = Reference(
        summary,
        min_col=1,
        min_row=header_row + 1,
        max_row=final_row,
    )
    for column in range(min_col, max_col + 1):
        values = Reference(
            summary,
            min_col=column,
            min_row=header_row,
            max_row=final_row,
        )
        chart.series.append(Series(values, day_values, title_from_data=True))
    chart.height = 7
    chart.width = 13
    return chart


def export_timecourse_xlsx(result: TimeCourseResult, path: str | Path) -> Path:
    """Save independently calculated time-point results in wide and long form."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    spec = mode_specification(result.matching_mode)
    database_provenance = file_provenance(result.database.source_path)
    sample_provenance = {}
    for point in result.timepoints:
        source = point.analysis.sample.source_path.resolve()
        if source not in sample_provenance:
            sample_provenance[source] = file_provenance(source)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Time Course Summary"
    summary["A1"] = APP_VERSION
    summary["A1"].font = Font(size=16, bold=True, color="17365D")
    summary["A2"] = result.series_name
    summary["A2"].font = Font(size=12, bold=True)
    summary["A3"] = "Each repertoire is analysed independently; frequencies are never pooled across days."

    headers = [
        "Day",
        "Label",
        "Sample ID",
        "Input format",
        "Repertoire file",
        "Listed reads",
        "Frequency denominator reads",
        "≤LV0 unique clones",
        "≤LV1 unique clones",
        "≤LV2 unique clones",
        "≤LV0 total reads",
        "≤LV1 total reads",
        "≤LV2 total reads",
        "≤LV0 frequency (%)",
        "≤LV1 frequency (%)",
        "≤LV2 frequency (%)",
        "LV0 unique clones",
        "LV1 unique clones",
        "LV2 unique clones",
        "LV0 total reads",
        "LV1 total reads",
        "LV2 total reads",
        "LV0 frequency (%)",
        "LV1 frequency (%)",
        "LV2 frequency (%)",
    ]
    header_row = 5
    summary.append([])
    summary.append(headers)
    _style_header(summary, header_row, 1, len(headers))
    for point in result.timepoints:
        analysis = point.analysis
        cumulative = analysis.cumulative_summaries
        exact = analysis.exact_summaries
        row = [
            point.spec.day,
            point.spec.label,
            analysis.sample.sample_id,
            analysis.sample.input_format,
            str(analysis.sample.source_path),
            analysis.sample.listed_reads,
            analysis.sample.total_reads,
            *[item.unique_clones for item in cumulative],
            *[item.total_reads for item in cumulative],
            *[item.frequency_percent for item in cumulative],
            *[item.unique_clones for item in exact],
            *[item.total_reads for item in exact],
            *[item.frequency_percent for item in exact],
        ]
        summary.append(row)
        for column in range(14, 17):
            summary.cell(summary.max_row, column).number_format = "0.000000000"
        for column in range(23, 26):
            summary.cell(summary.max_row, column).number_format = "0.000000000"

    if result.timepoints:
        final_row = header_row + len(result.timepoints)
        clone_chart = _timecourse_scatter_chart(
            summary,
            header_row=header_row,
            final_row=final_row,
            min_col=8,
            max_col=10,
            title="Cumulative unique clones",
            y_axis_title="Unique clones",
        )
        summary.add_chart(clone_chart, "AA5")

        frequency_chart = _timecourse_scatter_chart(
            summary,
            header_row=header_row,
            final_row=final_row,
            min_col=14,
            max_col=16,
            title="Cumulative frequency (%)",
            y_axis_title="Frequency (%)",
        )
        summary.add_chart(frequency_chart, "AA20")

    long_sheet = workbook.create_sheet("Long Summary")
    long_headers = [
        "Day",
        "Label",
        "Sample ID",
        "Aggregation",
        "Distance",
        "Unique clones",
        "Total reads",
        "Frequency (%)",
    ]
    long_sheet.append(long_headers)
    _style_header(long_sheet, 1, 1, len(long_headers))
    for point in result.timepoints:
        for aggregation, summaries in (
            ("Exact", point.analysis.exact_summaries),
            ("Cumulative", point.analysis.cumulative_summaries),
        ):
            for item in summaries:
                long_sheet.append(
                    [
                        point.spec.day,
                        point.spec.label,
                        point.analysis.sample.sample_id,
                        aggregation,
                        item.label,
                        item.unique_clones,
                        item.total_reads,
                        item.frequency_percent,
                    ]
                )
                long_sheet.cell(long_sheet.max_row, 8).number_format = "0.000000000"

    qc = workbook.create_sheet("Sample QC")
    qc_headers = [
        "Day",
        "Label",
        "Sample ID",
        "Input format",
        "Repertoire file",
        "Repertoire SHA-256",
        "Repertoire size (bytes)",
        "Repertoire modified time",
        "Source rows",
        "Accepted rows",
        "Skipped rows",
        "Normalized unique clones",
        "Listed reads",
        "Frequency denominator reads",
        "Metadata",
    ]
    qc.append(qc_headers)
    _style_header(qc, 1, 1, len(qc_headers))
    for point in result.timepoints:
        sample = point.analysis.sample
        provenance = sample_provenance[sample.source_path.resolve()]
        metadata = " | ".join(f"{key}: {value}" for key, value in sample.metadata.items())
        qc.append(
            [
                point.spec.day,
                point.spec.label,
                sample.sample_id,
                sample.input_format,
                str(provenance.path),
                provenance.sha256,
                provenance.size_bytes,
                provenance.modified_at,
                sample.source_rows,
                sample.accepted_rows,
                sample.skipped_rows,
                len(sample.clones),
                sample.listed_reads,
                sample.total_reads,
                metadata,
            ]
        )

    matched = workbook.create_sheet("Matched Clones")
    annotations = _annotation_order(result)
    matched_headers = [
        "Day",
        "Label",
        "Sample ID",
        "Distance",
        "IGHV normalized",
        "IGHJ normalized",
        "CDR3 normalized",
        "Reads",
        "Frequency (%)",
        "Database entries retained",
        "Database distances retained",
    ] + annotations
    matched.append(matched_headers)
    _style_header(matched, 1, 1, len(matched_headers))
    for point in result.timepoints:
        for match in point.analysis.matches:
            row = [
                point.spec.day,
                point.spec.label,
                point.analysis.sample.sample_id,
                f"LV{match.distance}",
                match.clone.display_v_gene,
                match.clone.display_j_gene,
                match.clone.cdr3_aa,
                match.clone.reads,
                match.clone.frequency_percent,
                len(match.database_entries),
                " | ".join(f"LV{distance}" for distance in match.entry_distances),
            ]
            row.extend(match.combined_annotation(column) for column in annotations)
            matched.append(row)
            matched.cell(matched.max_row, 9).number_format = "0.000000000"

    method = workbook.create_sheet("Method")
    method.append(["Item", "Fixed rule for this result"])
    _style_header(method, 1, 1, 2)
    method_rows = [
        ("Application version", APP_VERSION),
        ("Time-course module version", TIMECOURSE_VERSION),
        ("Matching algorithm version", ALGORITHM_VERSION),
        ("Analysis date", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Series name", result.series_name),
        ("Matching mode code", result.matching_mode.value),
        ("Matching mode", spec.label),
        ("Database input format", result.database.metadata.get("Database input format", "Unspecified")),
        ("Common database", str(result.database.source_path)),
        ("Common database SHA-256", database_provenance.sha256),
        ("Common database size (bytes)", database_provenance.size_bytes),
        ("Common database modified time", database_provenance.modified_at),
        ("Time-point order", "Ascending numeric Day; true numeric x spacing"),
        ("Allowed Day", "Any finite signed number, including negative, zero, positive, and decimal"),
        ("Duplicate Day", "Rejected; no automatic averaging or pooling"),
        ("Per-time-point calculation", "Each repertoire is loaded and matched independently"),
        ("Frequency calculation", "Calculated independently using each sample's own denominator"),
        ("Displayed classes", "Exact LV0/LV1/LV2 and cumulative ≤LV0/≤LV1/≤LV2"),
        ("Sample clone key", spec.sample_clone_key),
        ("Candidate selection", spec.candidate_rule),
        ("CDR3 handling", spec.cdr3_handling),
    ]
    method_rows.extend(runtime_provenance_rows())
    for row in method_rows:
        method.append(row)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet is not summary else "A6"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
        _autosize(sheet)
    workbook.save(destination)
    return destination
